import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from tolaris.dinheiro import formatar_brl, formatar_percentual
from tolaris.motor.caso import caso_de_json, caso_para_json
from tolaris.motor.modelos import AlteracaoSalarial, Aviso, DadosRescisao, Modalidade
from tolaris.motor.rescisao import calcular_rescisao
from tolaris.relatorios.excel import gerar_excel
from tolaris.relatorios.pdf import gerar_pdf

DADOS = DadosRescisao(
    admissao=date(2023, 3, 1),
    desligamento=date(2026, 5, 15),
    modalidade=Modalidade.SEM_JUSTA_CAUSA,
    aviso=Aviso.INDENIZADO,
    salario=Decimal("3000"),
    historico_salarial=[AlteracaoSalarial(date(2023, 3, 1), Decimal("2500"))],
    saldo_fgts=None,
    multa_467=True,
)


def test_formatacao_brasileira():
    assert formatar_brl(Decimal("1234567.891")) == "R$ 1.234.567,89"
    assert formatar_brl(Decimal("-10")) == "-R$ 10,00"
    assert formatar_percentual(Decimal("0.075")) == "7,5%"
    assert formatar_percentual(Decimal("0.14")) == "14%"


def test_caso_ida_e_volta():
    assert caso_de_json(caso_para_json(DADOS)) == DADOS


def test_pdf_e_gerado():
    conteudo = gerar_pdf(calcular_rescisao(DADOS))
    assert conteudo.startswith(b"%PDF")


def test_excel_tem_valores_numericos():
    resultado = calcular_rescisao(DADOS)
    planilha = load_workbook(io.BytesIO(gerar_excel(resultado))).active
    valores = [linha[3] for linha in planilha.iter_rows(values_only=True) if isinstance(linha[3], (int, float))]
    assert float(resultado.liquido) in valores
