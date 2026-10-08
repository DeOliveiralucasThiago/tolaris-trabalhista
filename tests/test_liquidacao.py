"""Liquidação de sentença: verbas deferidas, deduções, INSS, IR, honorários e custas.

Índices fictícios de tests/test_atualizacao.py (todos os meses: IPCA-E 1%, TR 0,1%, SELIC 1%,
IPCA 0,5%, taxa legal 0,5%). Valores esperados calculados à mão.
"""

from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest

from tests.test_atualizacao import INDICES
from tests.test_horas_extras import dados as dados_pedidos
from tests.test_relatorios_caso import DADOS as DADOS_RESCISAO
from tolaris.motor.liquidacao import CRITERIO_SELIC, calcular_liquidacao
from tolaris.motor.modelos import (
    DadosLiquidacao,
    ErroDeEntrada,
    Grupo,
    NaturezaPagamento,
    ValorPago,
    VerbaRescisoria,
)
from tolaris.motor.tributos import calcular_inss

# Horas extras de jan a mar/2022 (tests/test_horas_extras.py): 300,00 por mês + DSR, 13º, férias e aviso.
# Ajuizamento em 01/06/2022, liquidação em 10/03/2023.
HORAS_EXTRAS = dict(data_ajuizamento=date(2022, 6, 1), data_liquidacao=date(2023, 3, 10), pedidos=dados_pedidos())


def liquidar(**kwargs):
    return calcular_liquidacao(DadosLiquidacao(**{**HORAS_EXTRAS, **kwargs}), INDICES)


def inss_por_competencia(resultado):
    return {(linha.competencia, linha.decimo_terceiro): linha for linha in resultado.inss}


def test_inss_por_competencia_somando_o_salario_ja_pago():
    linhas = inss_por_competencia(liquidar())
    janeiro = linhas[(date(2022, 1, 1), False)]
    # Base deferida: horas 300,00 + DSR 72,00. Tabela de 2022 sobre 2.200 + 372 = 2.572:
    # 1.212 × 7,5% + 1.215,35 × 9% + 144,65 × 12% = 217,64; sobre 2.200: 179,82 → 37,82
    assert janeiro.base_devida == Decimal("372.00")
    assert janeiro.base_paga == Decimal("2200.00")
    assert janeiro.segurado == Decimal("37.82")
    assert janeiro.empresa == Decimal("103.42")  # 27,8% (20% + RAT 2% + terceiros 5,8%) × 372
    # Juros SELIC de março/2022 a fevereiro/2023 (12 × 1%) + 1% no mês da liquidação
    assert janeiro.criterio == CRITERIO_SELIC
    assert janeiro.fator == Decimal("0.13")
    # 13º: horas no 13º 100,00; 13º pago 2.200 × 3/12 = 550 → INSS(650) − INSS(550) = 7,50
    decimo = linhas[(date(2022, 3, 1), True)]
    assert (decimo.base_devida, decimo.base_paga, decimo.segurado) == (Decimal("100"), Decimal("550"), Decimal("7.50"))
    # Férias e aviso são indenizatórios: em março entram só horas 300,00 + DSR 44,44
    assert linhas[(date(2022, 3, 1), False)].base_devida == Decimal("344.44")


def test_inss_so_sobre_as_verbas_deferidas():
    linhas = inss_por_competencia(liquidar(considerar_salario_pago=False))
    assert linhas[(date(2022, 1, 1), False)].segurado == Decimal("27.90")  # 372 × 7,5%


def test_simples_nacional_sem_cota_patronal():
    assert liquidar(simples_nacional=True).inss_empresa == 0


def test_imposto_de_renda_acumulado_sem_juros():
    r = liquidar()
    # Tributável: horas 900 + DSR 166,44 + 13º 100 + correção (IPCA-E até maio; SELIC depois é juros):
    # horas 300 × (4,060401% + 3,0301% + 2,01%) = 27,30; DSR 5,33; 13º 2,01 → 1.201,08
    assert r.irrf.rendimento == Decimal("1201.08")
    assert r.irrf.ano_corrente is None
    assert r.irrf.meses == 4  # jan, fev e mar/2022 + 13º de 2022 (art. 37, § 1º, IN RFB nº 1.500/2014)
    assert r.irrf.valor == 0  # abaixo da isenção de 4 meses


def test_verbas_do_ano_da_liquidacao_pela_tabela_mensal():
    r = liquidar(data_liquidacao=date(2022, 8, 10))
    assert r.irrf.acumulado is None
    assert r.irrf.ano_corrente.rendimento == Decimal("1201.08")
    assert "art. 12-B" in r.irrf.memoria


def test_quadro_da_liquidacao_fecha():
    r = liquidar(honorarios_percentual=Decimal("0.10"))
    condenacao = r.bruto_atualizado + r.fgts_atualizado
    assert r.honorarios == (condenacao * Decimal("0.10")).quantize(Decimal("0.01"))
    assert r.custas == (condenacao * Decimal("0.02")).quantize(Decimal("0.01"))
    assert r.liquido_reclamante == r.bruto_atualizado - r.inss_segurado - r.irrf.valor
    assert r.total_reclamada == (
        r.bruto_atualizado + r.fgts_atualizado + r.inss_empresa + r.inss_acrescimos + r.honorarios + r.custas
    )
    rotulos = [linha[1] for linha in r.quadro]
    assert rotulos[0] == "Verbas deferidas atualizadas (sem FGTS)"
    assert rotulos[-1] == "Total devido pela reclamada"


def test_verbas_rescisorias_so_as_deferidas():
    deferidas = frozenset({VerbaRescisoria.SALDO_SALARIO, VerbaRescisoria.DECIMO_TERCEIRO, VerbaRescisoria.FGTS})
    r = calcular_liquidacao(
        DadosLiquidacao(date(2026, 6, 1), date(2026, 9, 10), rescisao=DADOS_RESCISAO, verbas_rescisorias=deferidas),
        INDICES,
    )
    assert {item.codigo for item in r.lancamentos} == {
        "saldo_salario",
        "decimo_terceiro_2026",
        "fgts_rescisorio",
        "multa_fgts",
    }
    assert not r.do_grupo(Grupo.DESCONTO)  # INSS e IR da rescisão são apurados de novo
    linhas = inss_por_competencia(r)
    # Saldo de salário (15 dias) e 13º não foram pagos: o salário pago do mês e do 13º é zero
    for decimo in (False, True):
        linha = linhas[(date(2026, 5, 1), decimo)]
        assert linha.base_paga == 0
        assert linha.segurado == calcular_inss(Decimal("1500"), date(2026, 5, 1)).valor == Decimal("112.50")


def test_valores_pagos_sao_deduzidos_com_fgts():
    pago = ValorPago("Horas extras pagas", date(2022, 2, 15), Decimal("100"), abater_fgts=True)
    r = liquidar(valores_pagos=[pago])
    valores = {item.codigo: item.valor for item in r.lancamentos}
    assert valores["pago_1"] == Decimal("-100")
    assert valores["pago_1_fgts"] == Decimal("-8.00")
    assert valores["pago_1_multa_fgts"] == Decimal("-3.20")
    assert inss_por_competencia(r)[(date(2022, 2, 1), False)].base_devida == Decimal("250.00")  # 300 + 50 − 100
    sem = liquidar()
    assert r.irrf.rendimento < sem.irrf.rendimento


def test_valor_pago_indenizatorio_nao_altera_inss():
    pago = ValorPago("Indenização", date(2022, 2, 1), Decimal("100"), NaturezaPagamento.INDENIZATORIA, True)
    r = liquidar(valores_pagos=[pago])
    assert inss_por_competencia(r)[(date(2022, 2, 1), False)].base_devida == Decimal("350.00")
    assert "pago_1_fgts" not in {item.codigo for item in r.lancamentos}


def test_honorarios_do_reclamante_e_justica_gratuita():
    base = dict(honorarios_reclamante_base=Decimal("10000"), honorarios_reclamante_percentual=Decimal("0.10"))
    suspensos = liquidar(justica_gratuita=True, **base)
    assert suspensos.honorarios_reclamante == Decimal("1000.00")
    assert suspensos.honorarios_reclamante_suspensos
    assert suspensos.liquido_reclamante == liquidar().liquido_reclamante
    assert any("ADI 5766" in alerta for alerta in suspensos.alertas)
    exigiveis = liquidar(justica_gratuita=False, **base)
    assert exigiveis.liquido_reclamante == liquidar().liquido_reclamante - Decimal("1000")
    assert exigiveis.total_reclamada == liquidar().total_reclamada  # retidos do crédito do reclamante


def test_custas_fixadas_na_sentenca():
    assert liquidar(custas_informadas=Decimal("500")).custas == Decimal("500.00")


def test_validacoes():
    with pytest.raises(ErroDeEntrada) as erro:
        calcular_liquidacao(DadosLiquidacao(date(2022, 6, 1), date(2022, 5, 1)), INDICES)
    mensagens = " ".join(erro.value.mensagens)
    assert "verbas deferidas" in mensagens
    assert "anterior ao ajuizamento" in mensagens


def test_relatorios_da_liquidacao():
    import io

    from openpyxl import load_workbook

    from tolaris.relatorios.excel import gerar_excel
    from tolaris.relatorios.pdf import gerar_pdf

    r = liquidar(honorarios_percentual=Decimal("0.10"), pedidos=replace(dados_pedidos()))
    assert gerar_pdf(r).startswith(b"%PDF")
    livro = load_workbook(io.BytesIO(gerar_excel(r)))
    assert livro.sheetnames[0] == "Liquidação"
    assert {"INSS", "Atualização", "Mês a mês"} <= set(livro.sheetnames)
    valores = [linha[1] for linha in livro["Liquidação"].iter_rows(values_only=True)]
    assert float(r.total_reclamada) in valores


def test_competencias_antes_de_marco_de_2009_sao_atualizadas_pelos_indices_trabalhistas():
    from tests.test_atualizacao import TAXAS
    from tolaris.datas import meses_entre
    from tolaris.motor.liquidacao import CRITERIO_TRABALHISTA
    from tolaris.tabelas import Indices

    tabela = Indices(
        {
            serie: {f"{mes:%Y-%m}": Decimal(taxa) for mes in meses_entre(date(2007, 1, 1), date(2026, 8, 31))}
            for serie, taxa in TAXAS.items()
        }
    )
    r = calcular_liquidacao(
        DadosLiquidacao(date(2010, 1, 15), date(2026, 3, 10), pedidos=dados_pedidos(ano=2009)), tabela
    )
    linhas = inss_por_competencia(r)
    janeiro = linhas[(date(2009, 1, 1), False)]
    # Jan/2009: 26 dias úteis e 5 repousos (01/01 + 4 domingos) → DSR 300 × 5/26 = 57,69; base 357,69.
    # Tabela de 03/2008 (alíquota única): 2.557,69 e 2.200 na faixa de 11% → 357,69 × 11% = 39,35
    assert janeiro.base_devida == Decimal("357.69")
    assert janeiro.segurado == Decimal("39.35")
    assert janeiro.criterio == CRITERIO_TRABALHISTA
    assert linhas[(date(2009, 3, 1), False)].criterio == CRITERIO_SELIC
