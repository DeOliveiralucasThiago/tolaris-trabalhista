"""Planilha com valores numéricos (não texto), para o advogado conferir e somar."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from tolaris.motor.atualizacao import ResultadoAtualizacao
from tolaris.motor.liquidacao import ResultadoLiquidacao
from tolaris.motor.modelos import Grupo, ResultadoCalculo
from tolaris.relatorios.mensal import colunas_ativas, numero

FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_HORAS = "#,##0.00"
AZUL = "002B5B"


def _cabecalho(planilha, titulos):
    planilha.append(titulos)
    for celula in planilha[planilha.max_row]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=AZUL)


def _aba_atualizacao(livro, atualizado: ResultadoAtualizacao):
    planilha = livro.create_sheet("Atualização")
    for linha in atualizado.criterio:
        planilha.append([linha])
    planilha.append([])
    _cabecalho(planilha, ["Rubrica", "Pedido", "Original", "Correção", "SELIC", "Juros", "Atualizado"])
    for item in atualizado.linhas + [None]:
        if item is None:
            valores, rotulo, pedido = atualizado.soma(), "Total (proventos e FGTS)", ""
        else:
            valores, rotulo, pedido = item.valores, item.descricao, item.pedido
        planilha.append(
            [rotulo, pedido, valores.original, valores.correcao, valores.selic, valores.juros, valores.total]
        )
        for coluna in range(3, 8):
            planilha.cell(planilha.max_row, coluna).number_format = FORMATO_MOEDA
    planilha.append(["Descontos (não atualizados)", "", None, None, None, None, atualizado.descontos])
    planilha.append(["Total atualizado", "", None, None, None, None, atualizado.total])
    for linha in (planilha.max_row - 1, planilha.max_row):
        planilha.cell(linha, 7).number_format = FORMATO_MOEDA
        planilha.cell(linha, 1).font = Font(bold=True)
    planilha.column_dimensions["A"].width = 60
    planilha.column_dimensions["B"].width = 28
    for coluna in "CDEFG":
        planilha.column_dimensions[coluna].width = 16


def _abas_liquidacao(livro, r: ResultadoLiquidacao):
    planilha = livro.create_sheet("Liquidação", 0)
    planilha.append(["TOLARIS TRABALHISTA – Resumo da liquidação"])
    planilha["A1"].font = Font(bold=True, size=14, color=AZUL)
    planilha.append([])
    secao_atual = None
    for secao, rotulo, valor, destaque in r.quadro:
        if secao != secao_atual:
            planilha.append([secao])
            planilha.cell(planilha.max_row, 1).font = Font(bold=True, color=AZUL)
            secao_atual = secao
        planilha.append([rotulo, valor])
        planilha.cell(planilha.max_row, 2).number_format = FORMATO_MOEDA
        if destaque:
            planilha.cell(planilha.max_row, 1).font = Font(bold=True)
            planilha.cell(planilha.max_row, 2).font = Font(bold=True)
    planilha.append([])
    planilha.append(["Imposto de renda", r.irrf.memoria])
    planilha.column_dimensions["A"].width = 60
    planilha.column_dimensions["B"].width = 20

    inss = livro.create_sheet("INSS")
    _cabecalho(
        inss,
        [
            "Competência",
            "13º",
            "Verbas deferidas",
            "Salário já pago",
            "Cota reclamante",
            "Cota reclamada",
            "Acréscimos",
            "Critério",
        ],
    )
    for linha in r.inss:
        inss.append(
            [
                linha.competencia,
                "Sim" if linha.decimo_terceiro else "",
                linha.base_devida,
                linha.base_paga,
                linha.segurado,
                linha.empresa,
                linha.acrescimos,
                linha.criterio,
            ]
        )
        inss.cell(inss.max_row, 1).number_format = "mm/yyyy"
        for coluna in range(3, 8):
            inss.cell(inss.max_row, coluna).number_format = FORMATO_MOEDA
    for coluna, largura in zip("ABCDEFGH", (14, 8, 18, 18, 18, 18, 16, 26), strict=True):
        inss.column_dimensions[coluna].width = largura


def gerar_excel(resultado: ResultadoCalculo, atualizado: ResultadoAtualizacao | None = None) -> bytes:
    livro = Workbook()
    planilha = livro.active
    planilha.title = "Resumo"

    planilha.append([f"TOLARIS TRABALHISTA – {resultado.titulo}"])
    planilha["A1"].font = Font(bold=True, size=14, color=AZUL)
    planilha.append([f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"])
    planilha.append([])

    for chave, valor in resultado.resumo.items():
        planilha.append([chave, valor])
    planilha.append([])

    _cabecalho(planilha, ["Grupo", "Rubrica", "Natureza", "Valor (R$)", "Memória de cálculo", "Fundamento"])
    for grupo in Grupo:
        for item in resultado.do_grupo(grupo):
            planilha.append(
                [item.grupo.value, item.descricao, item.natureza, item.valor, item.formula, item.fundamento]
            )
            planilha.cell(planilha.max_row, 4).number_format = FORMATO_MOEDA

    planilha.append([])
    for rotulo, valor, _destaque in resultado.totais():
        planilha.append(["", rotulo, "", valor])
        planilha.cell(planilha.max_row, 2).font = Font(bold=True)
        planilha.cell(planilha.max_row, 4).number_format = FORMATO_MOEDA
        planilha.cell(planilha.max_row, 4).font = Font(bold=True)

    if resultado.alertas:
        planilha.append([])
        planilha.append(["Alertas"])
        planilha.cell(planilha.max_row, 1).font = Font(bold=True)
        for alerta in resultado.alertas:
            planilha.append(["", alerta])

    for coluna, largura in zip("ABCDEF", (14, 55, 15, 15, 80, 70), strict=True):
        planilha.column_dimensions[coluna].width = largura
    for linha in planilha.iter_rows(min_col=5, max_col=6):
        for celula in linha:
            celula.alignment = Alignment(wrap_text=True, vertical="top")

    linhas = getattr(resultado, "mensal", None)
    if linhas:
        mensal = livro.create_sheet("Mês a mês")
        colunas = colunas_ativas(linhas)
        _cabecalho(mensal, [c.titulo for c in colunas])
        for linha in linhas:
            mensal.append([numero(c, linha) for c in colunas])
            for indice, coluna in enumerate(colunas, start=1):
                celula = mensal.cell(mensal.max_row, indice)
                if coluna.tipo == "moeda":
                    celula.number_format = FORMATO_MOEDA
                elif coluna.tipo == "horas":
                    celula.number_format = FORMATO_HORAS
                elif coluna.tipo == "mes":
                    celula.number_format = "mm/yyyy"
                elif coluna.tipo == "fracao":
                    celula.number_format = "0.00%"
        for indice in range(1, len(colunas) + 1):
            mensal.column_dimensions[mensal.cell(1, indice).column_letter].width = 16

    if isinstance(resultado, ResultadoLiquidacao):
        atualizado = atualizado or resultado.atualizado
        _abas_liquidacao(livro, resultado)
        livro.active = 0
    if atualizado:
        _aba_atualizacao(livro, atualizado)

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
